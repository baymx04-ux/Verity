import io

import pytest
from app.ingestion.parsers import DocumentParseError, parse_document
from docx import Document


def test_parses_plain_text() -> None:
    parsed = parse_document("notes.txt", "text/plain", b"Line one\nLine two")

    assert parsed.page_count == 1
    assert "Line two" in parsed.pages[0].text


def test_parses_markdown_sections() -> None:
    source = b"# Handbook\n\nIntro text\n\n## Leave\n\n25 days per year\n"

    parsed = parse_document("handbook.md", "text/markdown", source)

    sections = [page.section for page in parsed.pages]
    assert "Handbook" in sections
    assert "Leave" in sections
    assert any("25 days" in page.text for page in parsed.pages)


def test_parses_html_without_scripts() -> None:
    source = b"<html><body><h1>Title</h1><script>alert('x')</script><p>Body copy</p></body></html>"

    parsed = parse_document("page.html", "text/html", source)

    combined = " ".join(page.text for page in parsed.pages)
    assert "Body copy" in combined
    assert "alert" not in combined
    assert any(page.section == "Title" for page in parsed.pages)


def test_parses_docx_paragraphs_and_tables() -> None:
    document = Document()
    document.add_heading("Policy", level=1)
    document.add_paragraph("Employees get 25 days.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Region"
    table.rows[0].cells[1].text = "Days"
    buffer = io.BytesIO()
    document.save(buffer)

    parsed = parse_document(
        "policy.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        buffer.getvalue(),
    )

    combined = "\n".join(page.text for page in parsed.pages)
    assert "Employees get 25 days." in combined
    assert "Region | Days" in combined
    assert any(page.section == "Policy" for page in parsed.pages)


def test_rejects_unsupported_extension() -> None:
    with pytest.raises(DocumentParseError):
        parse_document("archive.zip", "application/zip", b"PK\x03\x04")


def test_rejects_empty_file() -> None:
    with pytest.raises(DocumentParseError):
        parse_document("empty.txt", "text/plain", b"")


def test_rejects_pdf_without_extractable_text() -> None:
    # A valid-enough PDF header that pypdf cannot extract text from.
    with pytest.raises(DocumentParseError):
        parse_document("scan.pdf", "application/pdf", b"%PDF-1.4 not really a pdf")
